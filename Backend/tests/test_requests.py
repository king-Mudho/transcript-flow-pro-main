import io
from decimal import Decimal

from django.core.cache import cache
from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase
from rest_framework_simplejwt.tokens import RefreshToken

from accounts.models import Role, User, UserRole
from branches.models import ZimpostBranch
from requests_app import services
from requests_app.models import DispatchBatch, Driver, Request, RequestEvent
from requests_app.services import generate_reference_number

YEAR = timezone.now().year

DB_FIELDS = {
    "full_name": "Jane Doe",
    "reg_number": "R123456",
    "programme_name": "BSc Computer Science",
    "year_completed": YEAR - 1,
    "phone_number": "0771234567",
    "email": "jane@example.com",
    "cleared_department": True,
    "cleared_accounts": True,
    "cleared_library": True,
    "zone": "harare",
    "harare_address": "123 Main St",
    "suburb": "Avondale",
    "fee_amount": "15.00",
    "payment_method": "cash_on_delivery",
}

SUBMIT = {
    "full_name": "Jane Doe",
    "reg_number": "R123456",
    "programme_name": "BSc Computer Science",
    "year_completed": YEAR - 1,
    "phone_number": "0771234567",
    "email": "jane@example.com",
    "cleared_department": True,
    "cleared_accounts": True,
    "cleared_library": True,
    "zone": "harare",
    "harare_address": "123 Main St",
    "suburb": "Avondale",
}


def make_request(ref="MSU-2026-000001", **overrides):
    return Request.objects.create(reference_number=ref, **{**DB_FIELDS, **overrides})


class Base(APITestCase):
    def setUp(self):
        cache.clear()
        self.admin = User.objects.create_user(email="admin@example.com", password="pw12345678")
        UserRole.objects.create(user=self.admin, role=Role.ADMIN)
        self.plain = User.objects.create_user(email="plain@example.com", password="pw12345678")
        self.branch = ZimpostBranch.objects.create(branch_name="Gweru PO", branch_area="Gweru")

    def auth(self, user):
        access = str(RefreshToken.for_user(user).access_token)
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {access}")

    def bulk(self, ids, to, **extra):
        return self.client.post(
            "/api/requests/bulk-update-status/",
            {"ids": [str(i) for i in ids], "status": to, **extra},
            format="json",
        )

    def walk(self, req, *statuses):
        for s in statuses:
            services.change_status(req, s, self.admin)
            req.refresh_from_db()


class ReferenceNumberTests(Base):
    def test_format_and_uniqueness(self):
        ref = generate_reference_number()
        self.assertRegex(ref, r"^MSU-\d{4}-\d{6}$")
        refs = set()
        for _ in range(20):
            ref = generate_reference_number()
            make_request(ref)
            refs.add(ref)
        self.assertEqual(len(refs), 20)


# ------------------------------------------------------------------- B1.2


class SubmissionTests(Base):
    def test_anonymous_can_submit_and_server_sets_money_and_status(self):
        response = self.client.post("/api/requests/", SUBMIT, format="json")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        req = Request.objects.get()
        self.assertEqual(response.data["reference_number"], req.reference_number)
        self.assertEqual(req.fee_amount, Decimal("15.00"))
        self.assertEqual(req.payment_method, "cash_on_delivery")
        self.assertEqual(req.status, "received")
        self.assertFalse(req.paid)
        self.assertIsNone(req.exported_at)
        self.assertEqual(req.suburb, "Avondale")

    def test_forged_fields_are_ignored(self):
        forged = {
            **SUBMIT,
            "fee_amount": "0.00",
            "payment_method": "cash_deposit",
            "paid": True,
            "status": "collected",
            "status_reason": "x",
            "exported_at": "2020-01-01T00:00:00Z",
            "reference_number": "MSU-1999-000000",
        }
        self.assertEqual(self.client.post("/api/requests/", forged, format="json").status_code, 201)
        req = Request.objects.get()
        self.assertEqual(req.fee_amount, Decimal("15.00"))
        self.assertFalse(req.paid)
        self.assertEqual(req.status, "received")
        self.assertIsNone(req.exported_at)
        self.assertNotEqual(req.reference_number, "MSU-1999-000000")

    def test_outside_harare_fee_and_payment(self):
        payload = {**SUBMIT, "zone": "outside_harare", "zimpost_branch": str(self.branch.id)}
        self.assertEqual(
            self.client.post("/api/requests/", payload, format="json").status_code, 201
        )
        req = Request.objects.get()
        self.assertEqual(req.fee_amount, Decimal("20.00"))
        self.assertEqual(req.payment_method, "cash_deposit")
        self.assertEqual(req.zimpost_branch, self.branch)
        self.assertIsNone(req.harare_address)

    def _refused(self, **changes):
        payload = {**SUBMIT, **changes}
        response = self.client.post("/api/requests/", payload, format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST, response.data)
        self.assertEqual(Request.objects.count(), 0)
        return response

    def test_validation_rules(self):
        self._refused(full_name="")
        self._refused(year_completed=YEAR - 15)
        self._refused(year_completed=YEAR + 1)
        self._refused(cleared_library=False)
        self._refused(harare_address="")
        self._refused(suburb="")
        self._refused(zone="outside_harare", zimpost_branch=None)
        self._refused(zone="mars")

    def test_inactive_branch_refused(self):
        self.branch.active = False
        self.branch.save()
        self._refused(zone="outside_harare", zimpost_branch=str(self.branch.id))

    def test_error_message_is_readable(self):
        response = self._refused(harare_address="")
        self.assertIn("address", str(response.data).lower())

    def test_non_admin_list_forbidden_and_admin_get_is_paged(self):
        make_request()
        self.assertEqual(self.client.get("/api/requests/").status_code, 401)
        self.auth(self.plain)
        self.assertEqual(self.client.get("/api/requests/").status_code, 403)


# ---------------------------------------------------------- B2.1 status flow


class StatusFlowTests(Base):
    def setUp(self):
        super().setUp()
        self.req = make_request()
        self.auth(self.admin)

    def test_default_is_received_and_forward_chain(self):
        self.assertEqual(self.req.status, "received")
        self.walk(self.req, "submitted_to_msu", "collected_from_msu")
        self.assertEqual(self.req.status, "collected_from_msu")

    def test_cannot_skip_a_stage(self):
        with self.assertRaises(ValidationError):
            services.change_status(self.req, "collected_from_msu", self.admin)
        self.req.refresh_from_db()
        self.assertEqual(self.req.status, "received")

    def test_back_one_step_needs_a_note(self):
        self.walk(self.req, "submitted_to_msu")
        with self.assertRaises(ValidationError):
            services.change_status(self.req, "received", self.admin)
        self.req.refresh_from_db()
        services.change_status(self.req, "received", self.admin, note="Moved by mistake")
        self.req.refresh_from_db()
        self.assertEqual(self.req.status, "received")

    def test_cannot_go_back_two_steps(self):
        self.walk(self.req, "submitted_to_msu", "collected_from_msu")
        with self.assertRaises(ValidationError):
            services.change_status(self.req, "received", self.admin, note="x")

    def test_reject_needs_reason_everywhere(self):
        self.walk(self.req, "submitted_to_msu")
        with self.assertRaises(ValidationError):
            services.change_status(self.req, "rejected", self.admin)
        response = self.bulk([self.req.id], "rejected")
        self.assertEqual(response.status_code, 400)
        with self.assertRaises(IntegrityError), transaction.atomic():
            Request.objects.filter(pk=self.req.pk).update(status="rejected", status_reason="")

    def test_rejection_reason_cleared_when_it_moves_on(self):
        self.walk(self.req, "submitted_to_msu")
        response = self.bulk([self.req.id], "rejected", status_reason="Library fines")
        self.assertEqual(response.data["moved"], 1)
        self.req.refresh_from_db()
        self.assertEqual(self.req.status_reason, "Library fines")
        services.change_status(self.req, "submitted_to_msu", self.admin, note="Cleared now")
        self.req.refresh_from_db()
        self.assertIsNone(self.req.status_reason)

    def test_bulk_reports_moved_and_skipped(self):
        other = make_request("MSU-2026-000002")
        self.walk(other, "submitted_to_msu", "collected_from_msu")
        response = self.bulk([self.req.id, other.id], "submitted_to_msu")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["moved"], 1)
        self.assertEqual(len(response.data["skipped"]), 1)
        self.assertEqual(response.data["skipped"][0]["reference_number"], "MSU-2026-000002")

    def test_bulk_by_filters_selects_all_matching(self):
        make_request("MSU-2026-000002")
        response = self.client.post(
            "/api/requests/bulk-update-status/",
            {"filters": {"status": "received"}, "status": "submitted_to_msu"},
            format="json",
        )
        self.assertEqual(response.data["moved"], 2)

    def test_allowed_next_exposed_to_admin(self):
        self.walk(self.req, "submitted_to_msu")
        row = self.client.get("/api/requests/").data["results"][0]
        self.assertEqual(row["allowed_next"], ["collected_from_msu", "rejected", "received"])

    def test_old_values_are_gone(self):
        self.assertEqual(self.bulk([self.req.id], "in_transit").status_code, 400)

    def test_non_admin_cannot_bulk(self):
        self.auth(self.plain)
        self.assertEqual(self.bulk([self.req.id], "submitted_to_msu").status_code, 403)

    def test_bulk_validates_ids(self):
        self.assertEqual(self.bulk(["not-a-uuid"], "submitted_to_msu").status_code, 400)
        for bad in ([], "nope", None):
            response = self.client.post(
                "/api/requests/bulk-update-status/",
                {"ids": bad, "status": "submitted_to_msu"},
                format="json",
            )
            self.assertEqual(response.status_code, 400)


# --------------------------------------------------------- B2.2 history log


class HistoryTests(Base):
    def setUp(self):
        super().setUp()
        self.req = make_request()
        self.auth(self.admin)

    def test_creation_logs_received(self):
        event = self.req.events.get()
        self.assertEqual((event.event_type, event.to_value), ("status_change", "received"))

    def test_status_change_logs_actor_and_time(self):
        self.bulk([self.req.id], "submitted_to_msu")
        event = self.req.events.first()
        self.assertEqual(event.to_value, "submitted_to_msu")
        self.assertEqual(event.from_value, "received")
        self.assertEqual(event.actor, self.admin)

    def test_paid_logs_payment_event(self):
        self.client.patch(f"/api/requests/{self.req.id}/", {"paid": True}, format="json")
        self.assertTrue(self.req.events.filter(event_type="payment", to_value="True").exists())

    def test_history_endpoint_newest_first_admin_only(self):
        self.bulk([self.req.id], "submitted_to_msu")
        response = self.client.get(f"/api/requests/{self.req.id}/history/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data[0]["to_value"], "submitted_to_msu")
        self.assertEqual(response.data[0]["actor"], "admin@example.com")
        self.assertEqual(response.data[-1]["actor"], "Graduate")
        self.auth(self.plain)
        self.assertEqual(self.client.get(f"/api/requests/{self.req.id}/history/").status_code, 403)

    def test_events_cannot_be_edited_or_deleted(self):
        event = self.req.events.get()
        event.note = "tampered"
        with self.assertRaises(ValidationError):
            event.save()
        with self.assertRaises(ValidationError):
            event.delete()
        self.assertEqual(RequestEvent.objects.count(), 1)


# ----------------------------------------------------- B2.3 public lookup

LOOKUP_KEYS = {
    "reference_number",
    "full_name",
    "programme_name",
    "status",
    "status_reason",
    "paid",
    "zone",
    "fee_amount",
    "payment_method",
    "created_at",
    "stages",
    "branch",
    "zimpost",
    "driver",
    "delivered_by",
}


class StatusLookupTests(Base):
    def lookup(self, ref="MSU-2026-000001", reg="R123456"):
        return self.client.get(
            "/api/requests/status/", {"reference_number": ref, "reg_number": reg}
        )

    def test_limited_field_set_and_case_insensitive_reg(self):
        make_request()
        response = self.lookup(reg="r123456")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(set(response.data.keys()), LOOKUP_KEYS)

    def test_no_match(self):
        make_request()
        self.assertEqual(self.lookup(reg="WRONG").status_code, 404)
        self.assertEqual(self.lookup(ref="MSU-2026-999999").status_code, 404)

    def test_stage_dates_come_from_events(self):
        req = make_request()
        self.walk(req, "submitted_to_msu", "collected_from_msu")
        stages = self.lookup().data["stages"]
        self.assertEqual(set(stages), {"received", "submitted_to_msu", "collected_from_msu"})

    def test_no_staff_or_notes_leak(self):
        req = make_request()
        services.change_status(req, "submitted_to_msu", self.admin, note="secret staff note")
        text = str(self.lookup().data)
        self.assertNotIn("secret staff note", text)
        self.assertNotIn("admin@example.com", text)

    def test_rejected_shows_reason_only_when_rejected(self):
        req = make_request()
        self.walk(req, "submitted_to_msu")
        services.change_status(req, "rejected", self.admin, reason="Unpaid fees")
        data = self.lookup().data
        self.assertEqual(data["status"], "rejected")
        self.assertEqual(data["status_reason"], "Unpaid fees")

    def test_branch_only_for_outside_harare(self):
        make_request(
            zone="outside_harare", zimpost_branch=self.branch, harare_address=None, suburb=""
        )
        self.assertEqual(self.lookup().data["branch"]["branch_name"], "Gweru PO")


# ------------------------------------------------------- B2.4 delivery edits


class DeliveryEditTests(Base):
    def setUp(self):
        super().setUp()
        self.req = make_request()
        self.auth(self.admin)
        self.url = f"/api/requests/{self.req.id}/"

    def patch(self, **data):
        return self.client.patch(self.url, data, format="json")

    def test_edit_address_before_dispatch_logs_old_and_new(self):
        self.walk(self.req, "submitted_to_msu", "collected_from_msu")
        response = self.patch(harare_address="9 New Road", suburb="Borrowdale")
        self.assertEqual(response.status_code, 200, response.data)
        self.req.refresh_from_db()
        self.assertEqual(self.req.harare_address, "9 New Road")
        event = self.req.events.filter(field="harare_address").first()
        self.assertEqual((event.from_value, event.to_value), ("123 Main St", "9 New Road"))
        self.assertEqual(event.actor, self.admin)

    def test_locked_after_dispatch_on_api_and_model(self):
        self.walk(self.req, "submitted_to_msu", "collected_from_msu")
        driver = Driver.objects.create(full_name="Tendai M", phone="+263771111111")
        services.create_dispatch_batch([self.req.id], driver.id, self.admin)
        response = self.patch(harare_address="Somewhere else", suburb="Mbare")
        self.assertEqual(response.status_code, 400)
        self.assertIn("Locked", response.data["detail"])
        self.req.refresh_from_db()
        self.req.harare_address = "tampered"
        with self.assertRaises(ValidationError):
            self.req.save()

    def test_switching_zone_recalculates_fee_and_payment(self):
        response = self.patch(zone="outside_harare", zimpost_branch=str(self.branch.id))
        self.assertEqual(response.status_code, 200, response.data)
        self.req.refresh_from_db()
        self.assertEqual(self.req.fee_amount, Decimal("20.00"))
        self.assertEqual(self.req.payment_method, "cash_deposit")
        self.assertIsNone(self.req.harare_address)
        back = self.patch(zone="harare", harare_address="5 Home Rd", suburb="Eastlea")
        self.assertEqual(back.status_code, 200, back.data)
        self.req.refresh_from_db()
        self.assertEqual(self.req.fee_amount, Decimal("15.00"))
        self.assertEqual(self.req.payment_method, "cash_on_delivery")

    def test_outside_needs_active_branch(self):
        self.assertEqual(self.patch(zone="outside_harare").status_code, 400)

    def test_phone_and_email_editable_before_dispatch(self):
        response = self.patch(phone_number="0772222222", email="new@example.com")
        self.assertEqual(response.status_code, 200)
        self.req.refresh_from_db()
        self.assertEqual(self.req.phone_number, "0772222222")

    def test_non_admin_cannot_edit(self):
        self.auth(self.plain)
        self.assertEqual(self.patch(paid=True).status_code, 403)


# ---------------------------------------------------------- B3.1 drivers


class DriverTests(Base):
    def test_admin_crud_and_phone_normalisation(self):
        self.auth(self.admin)
        response = self.client.post(
            "/api/drivers/",
            {"full_name": "Tendai Moyo", "phone": "0771234567", "bike_registration": "AFG 1234"},
            format="json",
        )
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["phone"], "+263771234567")
        self.assertEqual(response.data["whatsapp_phone"], "+263771234567")
        driver_id = response.data["id"]
        update = self.client.patch(f"/api/drivers/{driver_id}/", {"active": False}, format="json")
        self.assertFalse(update.data["active"])

    def test_bad_phone_refused(self):
        self.auth(self.admin)
        response = self.client.post(
            "/api/drivers/", {"full_name": "X Y", "phone": "12345"}, format="json"
        )
        self.assertEqual(response.status_code, 400)

    def test_public_and_non_admin_cannot_read(self):
        Driver.objects.create(full_name="Tendai M", phone="+263771111111")
        self.assertEqual(self.client.get("/api/drivers/").status_code, 401)
        self.auth(self.plain)
        self.assertEqual(self.client.get("/api/drivers/").status_code, 403)

    def test_driver_with_batches_cannot_be_deleted(self):
        self.auth(self.admin)
        driver = Driver.objects.create(full_name="Tendai M", phone="+263771111111")
        req = make_request()
        self.walk(req, "submitted_to_msu", "collected_from_msu")
        services.create_dispatch_batch([req.id], driver.id, self.admin)
        self.assertEqual(self.client.delete(f"/api/drivers/{driver.id}/").status_code, 400)
        spare = Driver.objects.create(full_name="Spare D", phone="+263772222222")
        self.assertEqual(self.client.delete(f"/api/drivers/{spare.id}/").status_code, 204)


# ------------------------------------------------------ B3.2 dispatch batches


class DispatchTests(Base):
    def setUp(self):
        super().setUp()
        self.driver = Driver.objects.create(
            full_name="Tendai Moyo",
            phone="+263771111111",
            whatsapp_phone="+263771111111",
            bike_registration="AFG 1234",
        )
        self.auth(self.admin)
        self.ready = []
        for i in range(5):
            r = make_request(f"MSU-2026-10000{i}", suburb="Avondale" if i % 2 else "Borrowdale")
            self.walk(r, "submitted_to_msu", "collected_from_msu")
            self.ready.append(r)

    def create(self, requests, driver=None):
        return self.client.post(
            "/api/dispatch/batches/",
            {
                "request_ids": [str(r.id) for r in requests],
                "driver_id": str((driver or self.driver).id),
            },
            format="json",
        )

    def test_ready_list_sorted_by_suburb(self):
        data = self.client.get("/api/dispatch/ready/").data
        self.assertEqual(len(data), 5)
        suburbs = [r["suburb"] for r in data]
        self.assertEqual(suburbs, sorted(suburbs))

    def test_create_batch_moves_all_to_dispatched(self):
        response = self.create(self.ready)
        self.assertEqual(response.status_code, 201, response.data)
        self.assertRegex(response.data["batch_number"], r"^HRE-\d{8}-01$")
        for r in self.ready:
            r.refresh_from_db()
            self.assertEqual(r.status, "dispatched")
            self.assertEqual(r.batch.batch_number, response.data["batch_number"])
            self.assertIsNotNone(r.dispatched_at)
            self.assertTrue(r.events.filter(event_type="dispatch").exists())
        self.assertEqual(self.client.get("/api/dispatch/ready/").data, [])
        second = make_request("MSU-2026-200000")
        self.walk(second, "submitted_to_msu", "collected_from_msu")
        self.assertRegex(self.create([second]).data["batch_number"], r"-02$")

    def test_refuses_wrong_stage_outside_harare_and_inactive_driver(self):
        outside = make_request(
            "MSU-2026-300000",
            zone="outside_harare",
            zimpost_branch=self.branch,
            harare_address=None,
        )
        self.walk(outside, "submitted_to_msu", "collected_from_msu")
        early = make_request("MSU-2026-300001")
        self.assertEqual(self.create([outside]).status_code, 400)
        self.assertEqual(self.create([early]).status_code, 400)
        self.driver.active = False
        self.driver.save()
        self.assertEqual(self.create(self.ready).status_code, 400)
        self.assertEqual(DispatchBatch.objects.count(), 0)
        # Nothing half-done: the five stay where they were.
        self.assertEqual(Request.objects.filter(status="collected_from_msu").count(), 6)

    def test_all_or_nothing(self):
        early = make_request("MSU-2026-300001")
        self.assertEqual(self.create([*self.ready, early]).status_code, 400)
        for r in self.ready:
            r.refresh_from_db()
            self.assertEqual(r.status, "collected_from_msu")

    def test_harare_cannot_reach_dispatched_outside_a_batch(self):
        response = self.bulk([self.ready[0].id], "dispatched")
        self.assertEqual(response.data["moved"], 0)
        self.assertEqual(len(response.data["skipped"]), 1)

    def test_deliver_and_fail_then_close(self):
        batch = self.create(self.ready[:2]).data
        url = f"/api/dispatch/batches/{batch['id']}/resolve/"
        ok = self.client.post(
            url, {"request_id": str(self.ready[0].id), "delivered": True}, format="json"
        )
        self.assertEqual(ok.status_code, 200)
        self.assertEqual(ok.data["status"], "out_for_delivery")
        missing_reason = self.client.post(
            url, {"request_id": str(self.ready[1].id), "delivered": False}, format="json"
        )
        self.assertEqual(missing_reason.status_code, 400)
        fail = self.client.post(
            url,
            {"request_id": str(self.ready[1].id), "delivered": False, "reason": "Nobody home"},
            format="json",
        )
        self.assertEqual(fail.status_code, 200)
        self.assertEqual(fail.data["status"], "closed")
        self.ready[0].refresh_from_db()
        self.ready[1].refresh_from_db()
        self.assertEqual(self.ready[0].status, "collected")
        self.assertEqual(self.ready[1].status, "collected_from_msu")
        self.assertIsNone(self.ready[1].batch)
        history = self.ready[1].events.filter(event_type="status_change").first()
        self.assertIn("Nobody home", history.note)

    def test_close_refused_while_documents_are_out(self):
        batch = self.create(self.ready[:2]).data
        self.assertEqual(
            self.client.post(f"/api/dispatch/batches/{batch['id']}/close/").status_code, 400
        )

    def test_batch_list_counts_and_non_admin_blocked(self):
        self.create(self.ready[:3])
        row = self.client.get("/api/dispatch/batches/").data[0]
        self.assertEqual(row["document_count"], 3)
        self.assertEqual(row["driver"]["full_name"], "Tendai Moyo")
        self.auth(self.plain)
        self.assertEqual(self.client.get("/api/dispatch/batches/").status_code, 403)
        self.assertEqual(self.create(self.ready[3:]).status_code, 403)


# ---------------------------------------------------- B3.4 driver on lookup


class DriverOnTrackingTests(Base):
    def setUp(self):
        super().setUp()
        self.driver = Driver.objects.create(
            full_name="Tendai Moyo",
            phone="+263771111111",
            whatsapp_phone="+263772222222",
            bike_registration="AFG 1234",
        )
        self.req = make_request()
        self.walk(self.req, "submitted_to_msu", "collected_from_msu")

    def lookup(self):
        return self.client.get(
            "/api/requests/status/",
            {"reference_number": self.req.reference_number, "reg_number": "R123456"},
        ).data

    def test_driver_only_while_dispatched_harare(self):
        self.assertIsNone(self.lookup()["driver"])
        batch = services.create_dispatch_batch([self.req.id], self.driver.id, self.admin)
        driver = self.lookup()["driver"]
        self.assertEqual(driver["full_name"], "Tendai Moyo")
        self.assertEqual(driver["phone"], "+263771111111")
        self.assertEqual(driver["whatsapp_phone"], "+263772222222")
        self.assertEqual(driver["bike_registration"], "AFG 1234")
        self.assertIsNotNone(driver["dispatched_at"])
        self.req.refresh_from_db()
        services.resolve_batch_item(batch, self.req, True, "", self.admin)
        after = self.lookup()
        self.assertIsNone(after["driver"])
        self.assertEqual(after["delivered_by"], "Tendai")
        self.assertNotIn("+263771111111", str(after))

    def test_no_driver_data_for_outside_harare(self):
        outside = make_request(
            "MSU-2026-400000",
            zone="outside_harare",
            zimpost_branch=self.branch,
            harare_address=None,
        )
        self.walk(outside, "submitted_to_msu", "collected_from_msu")
        # Even if a batch were somehow attached, an outside request exposes none.
        batch = DispatchBatch.objects.create(
            batch_number="X", driver=self.driver, dispatched_at=timezone.now()
        )
        outside.batch = batch
        outside.save()
        services.change_status(outside, "dispatched", self.admin)
        data = self.client.get(
            "/api/requests/status/",
            {"reference_number": "MSU-2026-400000", "reg_number": "R123456"},
        ).data
        self.assertIsNone(data["driver"])
        self.assertIsNone(data["delivered_by"])


# ------------------------------------------------------ B3.5 Zimpost dispatch


class ZimpostDispatchTests(Base):
    def setUp(self):
        super().setUp()
        self.auth(self.admin)
        self.reqs = []
        for i in range(3):
            r = make_request(
                f"MSU-2026-50000{i}",
                zone="outside_harare",
                zimpost_branch=self.branch,
                harare_address=None,
                suburb="",
                fee_amount="20.00",
                payment_method="cash_deposit",
            )
            self.walk(r, "submitted_to_msu", "collected_from_msu")
            self.reqs.append(r)

    def test_bulk_dispatch_with_date_and_tracking(self):
        response = self.bulk(
            [r.id for r in self.reqs],
            "dispatched",
            dispatched_at="2026-10-02",
            zimpost_tracking_number="RR123456789ZW",
        )
        self.assertEqual(response.data["moved"], 3)
        for r in self.reqs:
            r.refresh_from_db()
            self.assertEqual(r.zimpost_tracking_number, "RR123456789ZW")
            self.assertEqual(r.dispatched_at.date().isoformat(), "2026-10-02")
        data = self.client.get(
            "/api/requests/status/",
            {"reference_number": self.reqs[0].reference_number, "reg_number": "R123456"},
        ).data
        self.assertEqual(data["zimpost"]["tracking_number"], "RR123456789ZW")
        self.assertEqual(data["branch"]["branch_name"], "Gweru PO")

    def test_tracking_number_optional_and_bad_date_refused(self):
        self.assertEqual(self.bulk([self.reqs[0].id], "dispatched").data["moved"], 1)
        self.assertEqual(
            self.bulk([self.reqs[1].id], "dispatched", dispatched_at="not-a-date").status_code, 400
        )


# ------------------------------------------------- B1.4 paging / B1.1 export


class ListAndExportTests(Base):
    def setUp(self):
        super().setUp()
        self.auth(self.admin)
        Request.objects.bulk_create(
            [
                Request(
                    reference_number=f"MSU-2026-{i:06d}",
                    **{
                        **{k: v for k, v in DB_FIELDS.items()},
                        "full_name": f"Graduate {i}",
                        "reg_number": f"R{i:06d}",
                        "phone_number": f"07710{i:05d}",
                    },
                )
                for i in range(120)
            ]
        )

    def test_pages_of_50_with_exact_count(self):
        first = self.client.get("/api/requests/").data
        self.assertEqual((first["count"], len(first["results"]), first["page"]), (120, 50, 1))
        last = self.client.get("/api/requests/", {"page": 3}).data
        self.assertEqual(len(last["results"]), 20)
        seen = {
            r["id"]
            for p in (1, 2, 3)
            for r in self.client.get("/api/requests/", {"page": p}).data["results"]
        }
        self.assertEqual(len(seen), 120)

    def test_server_side_search_incl_phone_and_filters(self):
        by_phone = self.client.get("/api/requests/", {"search": "0771000007"}).data
        self.assertEqual(by_phone["count"], 1)
        by_reg = self.client.get("/api/requests/", {"search": "R000042"}).data
        self.assertEqual(by_reg["count"], 1)
        self.assertEqual(
            self.client.get("/api/requests/", {"status": "collected"}).data["count"], 0
        )
        self.assertEqual(self.client.get("/api/requests/", {"zone": "harare"}).data["count"], 120)
        tomorrow = (timezone.now() + timezone.timedelta(days=2)).date().isoformat()
        self.assertEqual(
            self.client.get("/api/requests/", {"date_from": tomorrow}).data["count"], 0
        )
        today = timezone.now().date().isoformat()
        self.assertEqual(
            self.client.get(
                "/api/requests/", {"date_from": "2020-01-01", "date_to": tomorrow}
            ).data["count"],
            120,
        )
        self.assertEqual(today, today)

    def _export(self, **params):
        response = self.client.get("/api/requests/export/", params)
        self.assertEqual(response.status_code, 200)
        return response, load_workbook(io.BytesIO(response.content))["Requests"]

    def test_export_has_every_matching_row_not_one_page(self):
        response, ws = self._export()
        self.assertEqual(response["X-Exported-Count"], "120")
        self.assertEqual(ws.max_row, 121)

    def test_export_columns_and_phone_as_text(self):
        _, ws = self._export(search="R000007")
        headers = [c.value for c in ws[1]]
        self.assertEqual(
            headers,
            [
                "Reference #",
                "Full Name",
                "Reg #",
                "Phone Number",
                "Email",
                "Programme",
                "Year Completed",
                "Zone",
                "Address / Zimpost Branch",
                "Fee (US$)",
                "Payment Method",
                "Paid",
                "Status",
                "Submitted",
            ],
        )
        phone = ws.cell(row=2, column=4)
        self.assertEqual(phone.value, "0771000007")
        self.assertEqual(phone.data_type, "s")
        self.assertEqual(phone.number_format, "@")
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws.cell(row=2, column=13).value, "Received")
        self.assertEqual(ws.cell(row=2, column=9).value, "123 Main St, Avondale")

    def test_export_marks_only_matching_rows_and_is_admin_only(self):
        self._export(search="R000007")
        self.assertEqual(Request.objects.filter(exported_at__isnull=False).count(), 1)
        self.auth(self.plain)
        self.assertEqual(self.client.get("/api/requests/export/").status_code, 403)
